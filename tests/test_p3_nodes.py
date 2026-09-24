"""Unit tests for the fieldwork pipeline nodes (CLAUDE.md build brief P3 §1-2):
discover, profile, plan, execute (+ G6 reconciliation, G7 contract
conformance), classify, find, prioritise (+ exposure de-duplication, §0.3),
act, export. Uses the "mini" test Skill fixture
(tests/fixtures/skills/mini/*.yaml, owned by the P2 primitives work) with tiny
CSV data this file writes itself into tmp_path -- never into tests/fixtures/,
which this session does not own."""

from __future__ import annotations

import dataclasses
from pathlib import Path

import pandas as pd
import pytest

from orchestrator import runs as runs_module
from orchestrator.adapters.export_storage import LocalExportStorage
from orchestrator.contract import ContractViolation, LocalFileDataSource
from orchestrator.nodes.context import NodeContext
from orchestrator.nodes.fieldwork import act, classify, discover, execute, export, find, prioritise, profile
from orchestrator.skills import load_skill
from tests.conftest import canonical_ts

MINI_SKILL_DIR = Path(__file__).parent / "fixtures" / "skills" / "mini"
AUDIT_PERIOD = ("2026-01-01", "2026-02-28")


def _fingerprint(fp_id: str, skill_content_hash: str | None = None) -> dict:
    return dict(
        fingerprint_id=fp_id,
        source_table_versions="{}",
        uploaded_file_hashes="{}",
        reference_data_hashes="{}",
        skill_content_hash=skill_content_hash,
        code_revision="rev1",
        dependency_lock_hash="dep1",
        runtime_config_hash="rc1",
        endpoint_config="{}",
        prompt_template_version="none",
        created_at=canonical_ts(0),
    )


def _write_mini_data(root: Path) -> None:
    claims = pd.DataFrame(
        [
            {"Employee ID": 1, "Transaction Date": "2026-01-05", "Amount": 100, "Vendor": "VendorA"},
            {"Employee ID": 1, "Transaction Date": "2026-01-06", "Amount": 600, "Vendor": "VendorA"},
            {"Employee ID": 2, "Transaction Date": "2026-01-10", "Amount": 700, "Vendor": "VendorB"},
            {"Employee ID": 3, "Transaction Date": "2026-01-15", "Amount": 50, "Vendor": "VendorC"},
            {"Employee ID": 4, "Transaction Date": "2026-02-01", "Amount": 900, "Vendor": "VendorD"},
        ]
    )
    register = pd.DataFrame(
        [
            {"Employee ID": 1, "Transaction Date": "2026-01-05", "Vendor": "VendorA"},
            {"Employee ID": 2, "Transaction Date": "2026-01-10", "Vendor": "VendorB"},
            {"Employee ID": 4, "Transaction Date": "2026-02-01", "Vendor": "VendorD"},
        ]
    )
    claims.to_csv(root / "claims.csv", index=False)
    register.to_csv(root / "register.csv", index=False)


@dataclasses.dataclass
class Harness:
    ctx: NodeContext
    state: object
    persistence: object


def _make_harness(local_persistence, tmp_path, *, run_owner="alice", engagement_id="ENG-DEFAULT", corrupt=None, options=None):
    persistence = local_persistence
    data_dir = tmp_path / "data"
    data_dir.mkdir()
    _write_mini_data(data_dir)
    if corrupt:
        corrupt(data_dir)

    skill = load_skill(MINI_SKILL_DIR)
    skill.validate()
    data_source = LocalFileDataSource(root_dir=data_dir, sources=skill.contract["sources"])
    source_versions = {name: data_source.resolve_version(name) for name in skill.contract["sources"]}

    fp = _fingerprint(f"FP-{tmp_path.name}", skill.content_hash)
    now = canonical_ts(1)
    state = runs_module.create_run(
        persistence, run_kind="fieldwork", engagement_id=engagement_id, skill_id=skill.skill_id,
        skill_version=skill.version, mode="playbook", audit_period=AUDIT_PERIOD, objective="t",
        run_owner=run_owner, options={"auto_confirm_plan": True, **(options or {})}, fingerprint=fp, now=now,
    )
    data_assets = [
        {"source": name, "table_fqn": name, "version": version}
        for name, version in source_versions.items()
    ]
    state = persistence.save_state(dataclasses.replace(state, data_assets=data_assets))

    export_dir = tmp_path / "exports"
    ctx = NodeContext(
        settings=type("S", (), {"catalog": None, "schema": None})(),
        persistence=persistence,
        data_source=data_source,
        skill=skill,
        clock=lambda: canonical_ts(2),
        export_storage=LocalExportStorage(root_dir=export_dir),
    )
    return Harness(ctx=ctx, state=state, persistence=persistence)


def _run_through_prioritise(h: Harness):
    state = discover(h.ctx, h.state)
    state = profile(h.ctx, state)
    state = execute(h.ctx, state)
    state = classify(h.ctx, state)
    state = find(h.ctx, state)
    state = prioritise(h.ctx, state)
    return state


# ── discover ──────────────────────────────────────────────────────────────────


def test_discover_confirms_bindings(local_persistence, tmp_path):
    h = _make_harness(local_persistence, tmp_path)
    state = discover(h.ctx, h.state)
    assert any(e["node"] == "discover" for e in state.events)


def test_discover_fails_on_missing_binding(local_persistence, tmp_path):
    h = _make_harness(local_persistence, tmp_path)
    stripped = dataclasses.replace(h.state, data_assets=[h.state.data_assets[0]])
    with pytest.raises(ContractViolation):
        discover(h.ctx, stripped)


def test_discover_fails_when_source_changed_since_pinning(local_persistence, tmp_path):
    h = _make_harness(local_persistence, tmp_path)
    tampered = [dict(b) for b in h.state.data_assets]
    for b in tampered:
        if b["source"] == "claims":
            b["version"] = "not-the-real-hash"
    state = dataclasses.replace(h.state, data_assets=tampered)
    with pytest.raises(ContractViolation) as exc:
        discover(h.ctx, state)
    assert "claims" in str(exc.value)


# ── profile ───────────────────────────────────────────────────────────────────


def test_profile_reports_row_and_null_counts(local_persistence, tmp_path):
    h = _make_harness(local_persistence, tmp_path)
    state = discover(h.ctx, h.state)
    state = profile(h.ctx, state)
    assert state.profile_result["claims"]["row_count"] == 5
    assert state.profile_result["claims"]["null_counts"]["Amount"] == 0
    assert state.profile_result["register"]["row_count"] == 3


# ── plan ──────────────────────────────────────────────────────────────────────


def test_plan_mirrors_skill_plan_tests(local_persistence, tmp_path):
    from orchestrator.nodes.fieldwork import plan as plan_node

    h = _make_harness(local_persistence, tmp_path)
    state = plan_node(h.ctx, h.state)
    test_ids = {t["test_id"] for t in state.plan["tests"]}
    assert test_ids == {"T1", "T2", "T3"}
    t3 = next(t for t in state.plan["tests"] if t["test_id"] == "T3")
    assert "not_testable" in t3


def test_plan_rejects_explorer_mode(local_persistence, tmp_path):
    from orchestrator.nodes.fieldwork import plan as plan_node

    h = _make_harness(local_persistence, tmp_path)
    explorer_state = dataclasses.replace(h.state, mode="explorer")
    with pytest.raises(ValueError):
        plan_node(h.ctx, explorer_state)


# ── execute + G6 reconciliation ─────────────────────────────────────────────


def test_execute_persists_metrics_and_flagged_rows(local_persistence, tmp_path):
    h = _make_harness(local_persistence, tmp_path)
    state = discover(h.ctx, h.state)
    state = execute(h.ctx, state)

    metrics = h.persistence.get_run_metrics(state.run_id)
    assert metrics["hv_count"]["value"] == 3
    assert metrics["hv_amount"]["value"] == 2200.0
    # mode: semi (mini/plan.yaml's T2) flags LEFT rows that DO have a right
    # match (a SQL semi-join), not unmatched ones -- claim1/claim3/claim5
    # each have a same-key register row; claim2/claim4 do not.
    assert metrics["missing_count"]["value"] == 3

    flagged = h.persistence.list_flagged_rows(state.run_id)
    flags = {r["flag"] for r in flagged}
    assert flags == {"RF_HV", "RF_MISSING"}
    assert sum(1 for r in flagged if r["flag"] == "RF_HV") == 3
    assert sum(1 for r in flagged if r["flag"] == "RF_MISSING") == 3

    # The mini Skill fixture (owned by the P2 primitives work) declares no
    # raw_<source> population, so the engine side of the G6 comparison is
    # unavailable here and variance is correctly None rather than falsely 0
    # (CLAUDE.md NN14 -- never claim a reconciliation that was not actually
    # computed). independent_rows (from DataSourceAdapter.row_count(), never
    # the engine's own frame) IS always available and correct.
    assert state.reconciliation["claims"]["independent_rows"] == 5
    assert state.reconciliation["claims"]["engine_rows"] is None
    assert state.reconciliation["claims"]["variance"] is None
    assert state.reconciliation["register"]["independent_rows"] == 3


def test_execute_never_touches_llm_state_fields(local_persistence, tmp_path):
    h = _make_harness(local_persistence, tmp_path)
    state = discover(h.ctx, h.state)
    state = execute(h.ctx, state)
    assert state.finding_narratives == {}
    assert state.exec_summary is None


# The mini Skill fixture declares no raw_<source> population (see the
# reconciliation assertion above), so a genuine G6 pass/fail comparison needs
# a Skill that does -- skills/tne_exco declares raw_<source> for every
# contract source. test_p3_tne_gates.py exercises G6 (pass and injected-
# failure) and G7 against that real Skill, using the small planted-exception
# fixture data (tests/fixtures/tne_planted/data/) rather than synthetic_data/
# so it runs in milliseconds, not minutes.


# ── G7 contract conformance ──────────────────────────────────────────────────


def test_g7_contract_violation_fails_the_run(local_persistence, tmp_path):
    def drop_amount_column(data_dir: Path) -> None:
        df = pd.read_csv(data_dir / "claims.csv")
        df = df.drop(columns=["Amount"])
        df.to_csv(data_dir / "claims.csv", index=False)

    h = _make_harness(local_persistence, tmp_path, corrupt=drop_amount_column)
    state = discover(h.ctx, h.state)
    with pytest.raises(ContractViolation) as exc:
        execute(h.ctx, state)
    assert "Amount" in str(exc.value)


# ── classify / find ───────────────────────────────────────────────────────────


def test_classify_summarises_exceptions(local_persistence, tmp_path):
    h = _make_harness(local_persistence, tmp_path)
    state = discover(h.ctx, h.state)
    state = execute(h.ctx, state)
    state = classify(h.ctx, state)
    by_id = {e["test_id"]: e for e in state.exceptions}
    assert by_id["T1"]["status"] == "exception"
    assert by_id["T1"]["exception_units"] == 3
    assert by_id["T3"]["status"] == "not_testable"


def test_find_builds_findings_from_persisted_metrics(local_persistence, tmp_path):
    h = _make_harness(local_persistence, tmp_path)
    state = discover(h.ctx, h.state)
    state = execute(h.ctx, state)
    state = classify(h.ctx, state)
    state = find(h.ctx, state)

    persisted = h.persistence.list_findings(state.run_id)
    by_test = {f["test_id"]: f for f in persisted}
    assert by_test["T1"]["severity"] == "High"
    assert by_test["T2"]["severity"] == "Low"
    assert {f["finding_id"] for f in state.findings} == {f["finding_id"] for f in persisted}

    issues = h.persistence.list_findings(state.run_id)  # sanity: still readable
    assert len(issues) == 2


def test_find_is_idempotent_on_reexecution(local_persistence, tmp_path):
    h = _make_harness(local_persistence, tmp_path)
    state = discover(h.ctx, h.state)
    state = execute(h.ctx, state)
    state = find(h.ctx, state)
    first = h.persistence.list_findings(state.run_id)

    state2 = find(h.ctx, state)
    second = h.persistence.list_findings(state2.run_id)
    assert [f["finding_id"] for f in first] == [f["finding_id"] for f in second]
    assert len(second) == 2  # never duplicated


# ── prioritise: exposure de-duplication (CLAUDE.md §0.3) ────────────────────


def test_prioritise_dedupes_exposure_across_findings(local_persistence, tmp_path):
    h = _make_harness(local_persistence, tmp_path)
    state = _run_through_prioritise(h)

    by_test = {f["test_id"]: f for f in state.findings}
    # RF_HV = claim2($600)+claim3($700)+claim5($900); RF_MISSING (mode: semi
    # flags rows that DO match the register) = claim1($100)+claim3($700)+
    # claim5($900). claim3 and claim5 are cited by BOTH findings -- the
    # naive sum-of-per-finding-totals would double count them.
    assert by_test["T1"]["exposure_amount"] == 2200.0  # claim2+claim3+claim5
    assert by_test["T2"]["exposure_amount"] == 1700.0  # claim1+claim3+claim5

    headline = h.persistence.get_run_metrics(state.run_id)["run_exposure_headline"]
    naive_sum = by_test["T1"]["exposure_amount"] + by_test["T2"]["exposure_amount"]
    assert naive_sum == 3900.0
    assert headline["value"] == 2300.0  # union {claim1,claim2,claim3,claim5}, not the naive sum
    assert headline["value"] < naive_sum


def test_prioritise_orders_by_severity_then_exposure(local_persistence, tmp_path):
    h = _make_harness(local_persistence, tmp_path)
    state = _run_through_prioritise(h)
    severities = [f["severity"] for f in state.findings]
    assert severities == sorted(severities, key=lambda s: {"High": 0, "Medium": 1, "Low": 2}[s])


# ── act ───────────────────────────────────────────────────────────────────────


def test_act_drafts_one_action_per_finding(local_persistence, tmp_path):
    h = _make_harness(local_persistence, tmp_path)
    state = _run_through_prioritise(h)
    state = act(h.ctx, state)

    actions = h.persistence.list_management_actions(filters={"run_id": state.run_id})
    assert len(actions) == 2
    assert all(a["status"] == "draft" for a in actions)
    assert {a["finding_id"] for a in actions} == {f["finding_id"] for f in state.findings}


def test_act_produces_no_actions_when_option_off(local_persistence, tmp_path):
    """CLAUDE.md §5 UI item 4 (NN13): "Generate management actions after
    review" unchecked must actually mean no actions get drafted, and the
    trace event must say why -- not just record the option and ignore it."""
    h = _make_harness(local_persistence, tmp_path, options={"generate_management_actions": False})
    state = _run_through_prioritise(h)
    state = act(h.ctx, state)

    actions = h.persistence.list_management_actions(filters={"run_id": state.run_id})
    assert actions == []
    assert state.management_actions == []
    assert any(
        e["node"] == "act" and e["message"] == "Management actions not generated (option off)"
        for e in state.events
    )


def test_act_reexecution_still_drafts_no_actions_when_option_off(local_persistence, tmp_path):
    """Idempotency (CLAUDE.md §2.3 rule 1): re-running act() with the option
    off must not leave stale actions from some other state around, and must
    not draft any either."""
    h = _make_harness(local_persistence, tmp_path, options={"generate_management_actions": False})
    state = _run_through_prioritise(h)
    state = act(h.ctx, state)
    state = act(h.ctx, state)

    actions = h.persistence.list_management_actions(filters={"run_id": state.run_id})
    assert actions == []


# ── export ────────────────────────────────────────────────────────────────────


def test_export_writes_xlsx_and_records_it(local_persistence, tmp_path):
    h = _make_harness(local_persistence, tmp_path)
    state = _run_through_prioritise(h)
    state = act(h.ctx, state)
    state = export(h.ctx, state)

    assert "xlsx" in state.exports
    path = Path(state.exports["xlsx"]["path"])
    assert path.is_file()
    assert path.stat().st_size > 0

    # `execute` (earlier in _run_through_prioritise) already recorded one
    # "frames:<source>" export per contract source (orchestrator/frames.py,
    # CLAUDE.md build brief P4 perf fix) -- `export` merges its own "xlsx"
    # entry into state.exports rather than replacing it (see export()'s own
    # docstring), so both kinds of export are recorded here, never just one.
    recorded = h.persistence.list_exports(state.run_id)
    xlsx_recorded = [e for e in recorded if e["kind"] == "xlsx"]
    frame_recorded = [e for e in recorded if e["kind"].startswith("frames:")]
    assert len(xlsx_recorded) == 1
    assert xlsx_recorded[0]["sha256"] == state.exports["xlsx"]["sha256"]
    assert {e["kind"] for e in frame_recorded} == {"frames:claims", "frames:register"}
    assert set(state.exports["frames"]) == {"claims", "register"}

    import hashlib

    assert hashlib.sha256(path.read_bytes()).hexdigest() == xlsx_recorded[0]["sha256"]


def test_export_produces_no_ticket_preview_when_option_off(local_persistence, tmp_path):
    """CLAUDE.md §5 UI item 4: "Prepare Jira ticket previews" unchecked
    (the default) means none at all -- not an empty preview -- and no
    Ticket Preview sheet in the workbook."""
    import openpyxl

    h = _make_harness(local_persistence, tmp_path)
    state = _run_through_prioritise(h)
    state = act(h.ctx, state)
    state = export(h.ctx, state)

    assert "ticket_preview" not in state.exports
    assert not any(e["kind"] == "ticket_preview" for e in h.persistence.list_exports(state.run_id))
    wb = openpyxl.load_workbook(Path(state.exports["xlsx"]["path"]))
    assert "Ticket Preview" not in wb.sheetnames


def test_export_produces_labelled_ticket_previews_when_option_on(local_persistence, tmp_path):
    """CLAUDE.md §5 UI item 4, §8: checked means one labelled 'Preview —
    not submitted' ticket per finding, stored with the run's other exports
    and included in the XLSX -- built through the tracker-neutral
    IssueTrackerAdapter (a ServiceNow/Jira/whichever-tracker implementation
    is a separate adapter later, never a change here)."""
    import openpyxl

    from orchestrator.adapters.issue_tracker_preview import TICKET_PREVIEW_STATUS

    h = _make_harness(local_persistence, tmp_path, options={"jira_preview_requested": True})
    state = _run_through_prioritise(h)
    state = act(h.ctx, state)
    state = export(h.ctx, state)

    assert "ticket_preview" in state.exports
    previews = state.exports["ticket_preview"]["ticket_previews"]
    assert len(previews) == len(state.findings)
    assert {p["issue_id"] for p in previews} == {f["finding_id"] for f in state.findings}
    assert all(p["status"] == TICKET_PREVIEW_STATUS for p in previews)

    recorded = [e for e in h.persistence.list_exports(state.run_id) if e["kind"] == "ticket_preview"]
    assert len(recorded) == 1
    assert recorded[0]["sha256"] == state.exports["ticket_preview"]["sha256"]

    wb = openpyxl.load_workbook(Path(state.exports["xlsx"]["path"]))
    assert "Ticket Preview" in wb.sheetnames
    sheet = wb["Ticket Preview"]
    assert sheet.cell(row=1, column=5).value == "status"
    assert sheet.cell(row=2, column=5).value == TICKET_PREVIEW_STATUS


def test_export_xlsx_never_writes_a_live_formula_cell(local_persistence, tmp_path):
    """P2/P3 gate review item 7 (OWASP CSV/formula-injection guidance): a
    data-derived string starting with =, +, -, @, tab or CR must be written
    as literal text, prefixed with a single quote, never as a live formula
    Excel would evaluate on open."""
    import openpyxl

    h = _make_harness(local_persistence, tmp_path)
    state = _run_through_prioritise(h)
    state = dataclasses.replace(state, objective="=HYPERLINK(\"http://evil.example\",\"click me\")")
    state = act(h.ctx, state)
    state = export(h.ctx, state)

    path = Path(state.exports["xlsx"]["path"])
    wb = openpyxl.load_workbook(path)
    cover = wb["Cover"]
    objective_cell = cover.cell(row=6, column=2)  # ("objective", value) row, value column
    assert objective_cell.value == "'=HYPERLINK(\"http://evil.example\",\"click me\")"
    assert objective_cell.data_type == "s"  # shared string, never "f" (formula)


def test_export_xlsx_cover_sheet_labels_a_self_approved_signoff(local_persistence, tmp_path):
    """CLAUDE.md §11 self sign-off decision: when signoff.approver equals
    run_owner, the exported workpaper's Cover sheet must carry the same
    "segregation of duties not enforced" statement shown in the UI -- an
    exported PPTX/XLSX that leaves this out is not a defensible workpaper."""
    import openpyxl

    from orchestrator.signoff_policy import SELF_APPROVED_LABEL

    h = _make_harness(local_persistence, tmp_path, run_owner="alice")
    state = _run_through_prioritise(h)
    state = act(h.ctx, state)
    state = dataclasses.replace(
        state,
        signoff={"approver": "alice", "timestamp": canonical_ts(9), "self_approved": True, "sod_enforced": False},
    )
    state = export(h.ctx, state)

    path = Path(state.exports["xlsx"]["path"])
    wb = openpyxl.load_workbook(path)
    cover = wb["Cover"]
    labels = [cover.cell(row=r, column=1).value for r in range(1, cover.max_row + 1)]
    values = [cover.cell(row=r, column=2).value for r in range(1, cover.max_row + 1)]
    assert "signed_off_by" in labels
    assert values[labels.index("signed_off_by")] == "alice"
    assert "signoff_note" in labels
    assert values[labels.index("signoff_note")] == SELF_APPROVED_LABEL


def test_export_xlsx_cover_sheet_omits_the_label_for_a_non_self_signoff(local_persistence, tmp_path):
    import openpyxl

    from orchestrator.signoff_policy import SELF_APPROVED_LABEL

    h = _make_harness(local_persistence, tmp_path, run_owner="alice")
    state = _run_through_prioritise(h)
    state = act(h.ctx, state)
    state = dataclasses.replace(
        state,
        signoff={"approver": "bob", "timestamp": canonical_ts(9), "self_approved": False, "sod_enforced": False},
    )
    state = export(h.ctx, state)

    path = Path(state.exports["xlsx"]["path"])
    wb = openpyxl.load_workbook(path)
    cover = wb["Cover"]
    labels = [cover.cell(row=r, column=1).value for r in range(1, cover.max_row + 1)]
    values = [cover.cell(row=r, column=2).value for r in range(1, cover.max_row + 1)]
    assert "signed_off_by" in labels
    assert values[labels.index("signed_off_by")] == "bob"
    assert "signoff_note" not in labels
    assert SELF_APPROVED_LABEL not in (v for v in values if v)


def test_export_xlsx_raises_if_severity_provenance_was_never_persisted(local_persistence, tmp_path):
    """CLAUDE.md §0.4/G8, item 3: never defaults analyst_set_severity to
    False -- a finding missing its persisted provenance must fail the export
    loudly rather than produce an unattributed severity in the workpaper."""
    h = _make_harness(local_persistence, tmp_path)
    state = _run_through_prioritise(h)
    state = act(h.ctx, state)

    findings = h.persistence.list_findings(state.run_id)
    assert findings, "fixture must produce at least one finding to corrupt"

    # write_findings' own contract requires analyst_set_severity/severity_basis
    # to be supplied by the caller -- simulate a caller bug (a future node that
    # forgets to carry them) by re-writing this run's findings without them.
    stripped = [dict(f) for f in findings]
    for f in stripped:
        f.pop("analyst_set_severity", None)
        f.pop("severity_basis", None)
    h.persistence.write_findings(
        state.run_id, stripped, engagement_id=state.engagement_id, skill_id=state.skill_id,
        skill_version=state.skill_version, now=canonical_ts(9),
    )

    from orchestrator.errors import MissingSeverityProvenance

    with pytest.raises(MissingSeverityProvenance):
        export(h.ctx, state)


def test_export_xlsx_never_writes_a_fabricated_zero_for_missing_amounts(local_persistence, tmp_path):
    """B4 (CLAUDE.md NN14, P2/P3 gate review): a non-monetary finding's
    exposure_amount is None, and a source whose raw_<source> population
    declares no amount_column has no reconciled amount at all -- neither may
    be written into the workpaper as a literal 0, which would misrepresent
    "not assessed in dollars"/"no amount column" as "assessed at zero"."""
    import openpyxl

    h = _make_harness(local_persistence, tmp_path)
    state = _run_through_prioritise(h)
    state = act(h.ctx, state)

    findings = h.persistence.list_findings(state.run_id)
    assert findings, "fixture must produce at least one finding to corrupt"
    # The mini fixture's own findings (T1, T2) are both monetary -- simulate a
    # non-monetary finding the way `prioritise` itself produces one, by
    # rewriting one finding's exposure fields to that exact shape.
    corrupted = [dict(f) for f in findings]
    corrupted[0]["exposure_amount"] = None
    corrupted[0]["exposure_basis"] = "non-monetary finding"
    h.persistence.write_findings(
        state.run_id, corrupted, engagement_id=state.engagement_id, skill_id=state.skill_id,
        skill_version=state.skill_version, now=canonical_ts(9),
    )

    state = export(h.ctx, state)
    path = Path(state.exports["xlsx"]["path"])
    wb = openpyxl.load_workbook(path)

    findings_ws = wb["Findings"]
    header = [c.value for c in findings_ws[1]]
    exposure_col = header.index("exposure_amount") + 1
    exposure_cell = findings_ws.cell(row=2, column=exposure_col)
    assert exposure_cell.value is None, "a non-monetary finding's exposure cell must be blank, never 0"

    # The mini fixture's plan.yaml declares no raw_<source> population for
    # either contract source, so both "claims" and "register" reconcile with
    # no declared amount column (orchestrator.populations.build_population:
    # amount is None when amount_column is not set) -- every row must read
    # the explicit n/a text, never a written 0.0.
    recon_ws = wb["Reconciliation"]
    recon_header = [c.value for c in recon_ws[1]]
    amount_col = recon_header.index("amount") + 1
    independent_amount_col = recon_header.index("independent_amount") + 1
    amount_variance_col = recon_header.index("amount_variance") + 1
    assert state.reconciliation
    data_rows = range(2, 2 + len(state.reconciliation))  # footer row follows, excluded
    for row in data_rows:
        for col in (amount_col, independent_amount_col, amount_variance_col):
            cell = recon_ws.cell(row=row, column=col)
            assert cell.value == "n/a — no amount column declared", (row, col, cell.value)
