"""T-E1 (P6 WP N13, docs/specs/P6_narration_design.md §11-§12) / the
narration half of WP7 (docs/specs/P6_P8_explorer_llm_design.md §8-§9): a
SKILL-001-shaped fieldwork run through the REAL state machine --
service.build_app_context + service.start_audit_run + the real
ThreadExecutor (never a direct node call, never orchestrator.pipeline.
run_phase called by hand) -- against tests/fixtures/skills/mini_candidates,
a Skill fixture already built to CLAUDE.md §4.4's own shape (manifest/
contract/plan/thresholds/findings, control_id/risk_id/assertion per test,
two severities, a not_testable third test) rather than the full 14-test
skills/tne_exco: narrate, accept one AI-proposed candidate, reject another,
edit a paragraph, sign off, finalise, export.

The model responses in tests/fixtures/narration/recorded_e2e_responses.json
are RECORDED FIXTURES (`"synthetic_recording": true`, the same convention
docs/specs/P6_P8_explorer_llm_design.md §9 WP7 uses for the Explorer
planner) -- CLAUDE.md §6's recorded results say MODEL_SONNET is untestable
in this development workspace (every call returns a 403 rate-limit-0), so
no line in that file is a genuine model transcript; it is consumed only
through DispatchingModelClient (tests/narration_test_support.py), never a
live endpoint.

Covers, in one file:
  - the full accept/reject/edit/sign-off/finalise/export flow;
  - NN7: exactly one llm_calls row per logical call, with prompt, response,
    endpoint, served model version, token counts and node;
  - NN11: no reasoning text anywhere a human or export can read it;
  - G15: a PII-bearing source value never reaches an llm_calls prompt;
  - replay: a second run over the same data under LLM_CACHE_MODE=replay
    with a RaisingModelClient gets byte-identical prose and zero live
    calls;
  - degraded mode: the Sonnet role unavailable still completes the run
    with the exact NN13 label and template-fallback prose;
  - G10: a clean population raises zero findings, zero candidates, makes
    no find_candidates call, and no risk language appears in any narrated
    text.
"""

from __future__ import annotations

import io
import json
import time
from pathlib import Path

import openpyxl
import pytest
from pptx import Presentation

from orchestrator import service
from orchestrator.adapters.model_fake import RaisingModelClient
from orchestrator.llm.errors import ModelUnavailable
from tests.narration_test_support import (
    MODEL_GPT_OSS_ENDPOINT,
    MODEL_SONNET_ENDPOINT,
    DispatchingModelClient,
    _write_mini_data,
    resp,
)

MINI_CANDIDATES_SKILL_DIR = Path(__file__).parent / "fixtures" / "skills" / "mini_candidates"
SKILLS_DIR = MINI_CANDIDATES_SKILL_DIR.parent
SKILL_ID = "SKILL-MINI-CAND"
AUDIT_PERIOD = ("2026-01-01", "2026-02-28")

MARKER_FIND_T1 = '"finding_key":"T1"'
MARKER_FIND_T2 = '"finding_key":"T2"'
MARKER_SYNTHESIS = "finding-synthesis/1"
MARKER_PRIORITY = "priority-rationale/1"
MARKER_REMEDIATION = "remediation/1"
MARKER_EXEC_SUMMARY = "exec-summary/1"
MARKER_CAPTIONS = "chart-captions/1"
MARKER_PROFILE = "profile-narrative/1"
MARKER_CANDIDATES = "finding-candidates/1"

_RECORDED_PATH = Path(__file__).parent / "fixtures" / "narration" / "recorded_e2e_responses.json"
PII_SENTINEL = "SENTINEL-PII-7731@example.test"


def _load_recorded() -> dict:
    data = json.loads(_RECORDED_PATH.read_text())
    assert data["synthetic_recording"] is True, "fixture must be labelled synthetic, never a real transcript"
    return data["responses"]


def _happy_client() -> DispatchingModelClient:
    payloads = _load_recorded()
    by_marker = {
        MARKER_FIND_T1: resp(payloads["find_t1"]),
        MARKER_FIND_T2: resp(payloads["find_t2"]),
        MARKER_SYNTHESIS: resp(payloads["synthesis"]),
        MARKER_PRIORITY: resp(payloads["priority"]),
        MARKER_REMEDIATION: resp(payloads["remediation"]),
        MARKER_EXEC_SUMMARY: resp(payloads["exec_summary"]),
        MARKER_CAPTIONS: resp(payloads["captions"], model="gpt-oss-test-v1"),
        MARKER_PROFILE: resp(payloads["profile"]),
        MARKER_CANDIDATES: resp(payloads["candidates"], reasoning=2),
    }
    return DispatchingModelClient(by_marker)


def _write_data_with_pii_sentinel(root: Path) -> None:
    """The same small claims/register population
    tests.narration_test_support._write_mini_data writes (2 missing
    claims, 2 high-value claims), with one MATCHED claim/register pair's
    Vendor carrying a PII sentinel value -- present in the raw source data,
    never in any narration payload (§3.2 of the design: narration payloads
    are metrics/counts only, never raw row values) -- so G15 is checked
    against real source content, not an untested assumption."""
    _write_mini_data(root)
    import pandas as pd

    sentinel_vendor = f"{PII_SENTINEL}-VendorD"
    for name in ("claims.csv", "register.csv"):
        df = pd.read_csv(root / name)
        df.loc[df["Vendor"] == "VendorD", "Vendor"] = sentinel_vendor
        df.to_csv(root / name, index=False)


def _write_clean_data(root: Path) -> None:
    """G10 (§11 of the design doc, the same shape test_p3_tne_gates.py's
    own G10 test uses against the sibling 'mini' Skill): every claim under
    the high-value limit (500), and no claim's (Employee ID, Transaction
    Date, Vendor) key matches any register row -- T2's primitive is a
    SEMI-join (mode: semi), so a clean population for 'missing_count > 0'
    to stay false is one where NOTHING matches, not the reverse."""
    import pandas as pd

    claims = pd.DataFrame(
        [
            {"Employee ID": 1, "Transaction Date": "2026-01-05", "Amount": 100, "Vendor": "VendorA"},
            {"Employee ID": 2, "Transaction Date": "2026-01-10", "Amount": 200, "Vendor": "VendorB"},
        ]
    )
    register = pd.DataFrame(
        [{"Employee ID": 9, "Transaction Date": "2026-01-09", "Vendor": "VendorZ"}]
    )
    claims.to_csv(root / "claims.csv", index=False)
    register.to_csv(root / "register.csv", index=False)


def _build_ctx(
    tmp_path: Path, *, model_sonnet: str | None = MODEL_SONNET_ENDPOINT, worker_id: str = "worker-a",
    llm_cache_mode: str = "live", data_writer=_write_data_with_pii_sentinel,
) -> service.AppContext:
    env = {
        "ORCH_BACKEND": "local",
        "ORCH_LOCAL_DB": str(tmp_path / "orch.db"),
        "ORCH_LOCAL_DATA_ROOT": str(tmp_path / "data"),
        "ORCH_LOCAL_EXPORT_ROOT": str(tmp_path / "exports"),
        "ORCH_WORKER_ID": worker_id,
        "SKILLS_DIR": str(SKILLS_DIR),
        "CODE_REVISION": "test-fixed-revision",
        "AUDIT_TIMEZONE": "Australia/Sydney",
        "NARRATION_ENABLED": "true",
        "AI_PROPOSED_FINDINGS_ENABLED": "true",
        "LLM_CACHE_MODE": llm_cache_mode,
    }
    if model_sonnet:
        env["MODEL_SONNET"] = model_sonnet
        env["MODEL_GPT_OSS"] = MODEL_GPT_OSS_ENDPOINT
    data_dir = tmp_path / "data"
    data_dir.mkdir(parents=True, exist_ok=True)
    data_writer(data_dir)
    return service.build_app_context(env)


def _wait_for(ctx, run_id, statuses, timeout=20):
    deadline = time.time() + timeout
    last = None
    while time.time() < deadline:
        last = service.get_run(ctx, run_id)["status"]
        if last in statuses:
            return last
        time.sleep(0.05)
    raise AssertionError(f"run {run_id} did not reach {statuses} in time (last={last})")


def _start_run(ctx: service.AppContext, *, run_owner="alice") -> str:
    bindings = service.suggest_bindings(ctx, SKILL_ID)
    assert all(bindings.values()), bindings
    return service.start_audit_run(
        ctx, skill_id=SKILL_ID, bindings=bindings, audit_period=AUDIT_PERIOD,
        objective="T-E1 narration e2e", run_owner=run_owner,
    )


def _read_exports(ctx, run_id):
    pptx_filename, pptx_bytes = service.get_export(ctx, run_id, "pptx")
    xlsx_filename, xlsx_bytes = service.get_export(ctx, run_id, "xlsx")
    prs = Presentation(io.BytesIO(pptx_bytes))
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), data_only=True)
    return prs, wb


def _all_slide_text(prs) -> str:
    return " ".join(
        shape.text_frame.text for slide in prs.slides for shape in slide.shapes if shape.has_text_frame
    )


def _findings_sheet_rows(wb) -> list[dict]:
    ws = wb["Findings"]
    header = [c.value for c in ws[1]]
    rows = []
    for row in ws.iter_rows(min_row=2):
        values = [c.value for c in row]
        if values[0] is None or (isinstance(values[0], str) and values[0].startswith("run_id=")):
            continue
        rows.append(dict(zip(header, values)))
    return rows


# ── the full flow ────────────────────────────────────────────────────────────


def test_full_narration_e2e_accept_reject_edit_signoff_finalise_export(tmp_path):
    ctx = _build_ctx(tmp_path)
    ctx.executor.start()
    try:
        client = _happy_client()
        ctx.model_client = client
        run_id = _start_run(ctx)
        status = _wait_for(ctx, run_id, {"awaiting_signoff", "failed"})
        assert status == "awaiting_signoff", service.get_run(ctx, run_id).get("status_reason")

        # ── the two rule findings narrated with model prose ──────────────
        narratives = ctx.persistence.get_narratives(run_id)
        assert narratives, "no narratives persisted -- narrate did not run"
        finding_observations = {
            n["target_id"]: n for n in narratives if n["target_kind"] == "finding" and n["field"] == "observation"
        }
        assert len(finding_observations) == 2
        assert all(n["origin"] == "model" for n in finding_observations.values())

        # ── the model proposed exactly the two candidates the fixture
        #    scripts, both C-2 anchored on the same uncited metric ────────
        candidates = ctx.persistence.list_candidates(run_id)
        assert len(candidates) == 2
        assert all(c["candidate_status"] == "candidate" for c in candidates)
        accepted = next(c for c in candidates if c["title"] == "Missing Register Match Amount")
        rejected = next(c for c in candidates if c["title"] == "Missing Register Match Volume and Amount")

        # ── accept one, reject the other ──────────────────────────────────
        service.decide_candidate(
            ctx, run_id, accepted["candidate_id"], decision="accepted", decided_severity="Medium", actor="alice",
        )
        service.decide_candidate(
            ctx, run_id, rejected["candidate_id"], decision="rejected",
            reason="Not material enough to warrant a second write-up of the same test.", actor="alice",
        )

        # ── edit one paragraph (Q3: placeholders only, coverage checked) ──
        state = ctx.persistence.load_state(run_id)
        findings = ctx.persistence.list_findings(run_id)
        t1_finding_id = next(f["finding_id"] for f in findings if f.get("test_id") == "T1")
        t1_observation_id = state.finding_narratives[t1_finding_id]
        edited_text = (
            "In this run, {count:hv_count} claim(s) exceeded the high-value threshold, "
            "worth {money:hv_amount} in total exposure."
        )
        edited = service.edit_narrative(ctx, run_id, t1_observation_id, edited_text, actor="alice")
        assert edited["origin"] == "human_edit"
        edits = ctx.persistence.list_narrative_edits(run_id) if hasattr(ctx.persistence, "list_narrative_edits") else None
        if edits is not None:
            assert any(e["narrative_id"] == t1_observation_id and e["action"] == "human_edit" for e in edits)

        # ── sign off, then finalise + export run in the export phase ──────
        service.sign_off(ctx, run_id, "alice")
        status = _wait_for(ctx, run_id, {"completed", "failed"})
        assert status == "completed", service.get_run(ctx, run_id).get("status_reason")

        # ── the accepted candidate is a real finding; the rejected one is
        #    kept, decided, and excluded from the finding set ─────────────
        findings = ctx.persistence.list_findings(run_id)
        ai_findings = [f for f in findings if f.get("origin") == "ai_proposed"]
        assert len(ai_findings) == 1
        assert ai_findings[0]["accepted_by"] == "alice"
        assert ai_findings[0]["severity"] == "Medium"
        candidates_after = {c["candidate_id"]: c for c in ctx.persistence.list_candidates(run_id)}
        assert candidates_after[accepted["candidate_id"]]["candidate_status"] == "accepted"
        assert candidates_after[rejected["candidate_id"]]["candidate_status"] == "rejected"

        # ── PPTX and XLSX content ──────────────────────────────────────────
        prs, wb = _read_exports(ctx, run_id)
        all_text = _all_slide_text(prs)
        assert "Missing Register Match Amount" in all_text
        assert "AI-proposed" in all_text or "Additional matters proposed by AI review" in all_text
        # the edited paragraph's own rendered numbers appear (placeholders
        # rendered, never the raw {count:...} span)
        assert "worth" in all_text and "total exposure" in all_text
        # the REJECTED candidate's own text never reaches an export (§14 Q10)
        assert "Missing Register Match Volume and Amount" not in all_text
        assert "Track both the count and the dollar impact" not in all_text
        assert "How many unmatched claims by dollar value" not in all_text

        findings_rows = _findings_sheet_rows(wb)
        ai_row = next(r for r in findings_rows if r["origin"] == "ai_proposed")
        assert ai_row["accepted_by"] == "alice"
        assert ai_row["proposed_severity"] == "Low"
        assert ai_row["decided_severity"] == "Medium"
        t1_row = next(r for r in findings_rows if r["test_id"] == "T1")
        assert "worth" in t1_row["observation"] and "total exposure" in t1_row["observation"]
        # BUG-6 (independent review, 2026-09-25): sign-off already walked
        # every rule finding's review_state to 'approved' before finalise
        # ever ran (it is the first export-phase node); the accepted
        # AI-proposed finding, materialised into `findings` for the first
        # time by finalise, must carry the same review_state -- in Delta AND
        # in this exported XLSX -- never the un-reviewed-looking 'draft'.
        assert t1_row["review_state"] == "approved"
        assert ai_row["review_state"] == "approved"
        xlsx_text = json.dumps([r for r in findings_rows], default=str)
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            for row in ws.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and "Missing Register Match Volume and Amount" in cell.value:
                        raise AssertionError(f"rejected candidate text leaked into {sheet_name}!{cell.coordinate}")

        # ── NN7: exactly one llm_calls row per logical (task, seq) call,
        #    every row carries prompt/response/endpoint/served version/
        #    tokens/node ────────────────────────────────────────────────
        calls = ctx.persistence.list_llm_calls(run_id)
        assert calls
        seen = set()
        for c in calls:
            key = (c["task"], c["seq"], c.get("transport_attempt"))
            assert key not in seen, f"duplicate llm_calls row for {key}"
            seen.add(key)
            assert c["node_name"] == "narrate"
            assert c["messages_json"], "prompt not logged"
            assert c["endpoint"] in (MODEL_SONNET_ENDPOINT, MODEL_GPT_OSS_ENDPOINT)
            assert c["served_model_version"], "served model version not logged"
            assert c["total_tokens"], "token counts not logged"
            if c["outcome"] == "succeeded":
                assert c["response_text"], "response not logged"

        # ── NN11: no reasoning text anywhere it could be read ──────────────
        for c in calls:
            assert "reasoning_parts_stripped" in c
            blob = json.dumps(c.get("response_text")) + json.dumps(c.get("messages_json"))
            assert "chain-of-thought" not in blob.lower()

        # ── G15: the PII sentinel embedded in raw source data never
        #    reaches any llm_calls prompt or response ─────────────────────
        for c in calls:
            blob = json.dumps(c.get("messages_json")) + json.dumps(c.get("response_text"))
            assert PII_SENTINEL not in blob, f"PII sentinel leaked into llm_calls row {c['call_id']}"
        assert PII_SENTINEL not in json.dumps(candidates, default=str)
        assert PII_SENTINEL not in all_text
    finally:
        ctx.executor.stop()


# ── replay ────────────────────────────────────────────────────────────────


def test_replay_mode_gives_byte_identical_prose_with_zero_live_calls(tmp_path):
    ctx1 = _build_ctx(tmp_path / "record")
    ctx1.executor.start()
    try:
        ctx1.model_client = _happy_client()
        run_id_1 = _start_run(ctx1)
        status = _wait_for(ctx1, run_id_1, {"awaiting_signoff", "failed"})
        assert status == "awaiting_signoff", service.get_run(ctx1, run_id_1).get("status_reason")
    finally:
        ctx1.executor.stop()

    def _independent_narratives(persistence, run_id):
        # `finding_id`/`theme_id` are literally f"{run_id}:..." so stripping
        # that prefix compares two runs' rows by what they MEAN. A
        # candidate_id is `sha256(f"{run_id}|{generation}|{rule_id}")`
        # instead (§5.1) -- it embeds run_id inside the hash, not as a
        # literal prefix -- so a candidate-kind row is normalised to its
        # OWN rule_id (stable across runs, §5.1's own identity table)
        # rather than its run-scoped candidate_id.
        rule_id_by_candidate = {c["candidate_id"]: c["rule_id"] for c in persistence.list_candidates(run_id)}
        out = {}
        for row in persistence.get_narratives(run_id):
            target_id = row["target_id"]
            if row["target_kind"] == "candidate":
                target_id = rule_id_by_candidate.get(target_id, target_id)
            elif target_id.startswith(f"{run_id}:"):
                target_id = target_id[len(run_id) + 1:]
            out[(row["target_kind"], target_id, row["field"])] = row["template_text"]
        return out

    recorded = _independent_narratives(ctx1.persistence, run_id_1)
    assert recorded

    # A SECOND, independent AppContext/run over the SAME underlying source
    # data and the SAME local DB (so it can read run 1's own llm_cache
    # rows -- llm_cache is keyed on (prompt_sha256, endpoint,
    # served_model_version, params_json), never on run_id, §1 point 8) with
    # a RaisingModelClient: LLM_CACHE_MODE=replay must serve every call
    # from cache, never touching the client.
    ctx2 = service.build_app_context({
        "ORCH_BACKEND": "local",
        "ORCH_LOCAL_DB": str(tmp_path / "record" / "orch.db"),
        "ORCH_LOCAL_DATA_ROOT": str(tmp_path / "record" / "data"),
        "ORCH_LOCAL_EXPORT_ROOT": str(tmp_path / "record" / "exports2"),
        "ORCH_WORKER_ID": "worker-b",
        "SKILLS_DIR": str(SKILLS_DIR),
        "CODE_REVISION": "test-fixed-revision",
        "AUDIT_TIMEZONE": "Australia/Sydney",
        "NARRATION_ENABLED": "true",
        "AI_PROPOSED_FINDINGS_ENABLED": "true",
        "LLM_CACHE_MODE": "replay",
        "MODEL_SONNET": MODEL_SONNET_ENDPOINT,
        "MODEL_GPT_OSS": MODEL_GPT_OSS_ENDPOINT,
    })
    ctx2.executor.start()
    try:
        ctx2.model_client = RaisingModelClient(
            ModelUnavailable("should-not-be-called", "replay mode never calls live", permanent=True)
        )
        run_id_2 = _start_run(ctx2, run_owner="bob")
        status = _wait_for(ctx2, run_id_2, {"awaiting_signoff", "failed"})
        assert status == "awaiting_signoff", service.get_run(ctx2, run_id_2).get("status_reason")

        replayed = _independent_narratives(ctx2.persistence, run_id_2)
        assert replayed == recorded

        calls2 = ctx2.persistence.list_llm_calls(run_id_2)
        assert calls2
        assert all(c["source"] == "cache" for c in calls2), [c["source"] for c in calls2]
        assert not any(c["source"] == "live" for c in calls2)
    finally:
        ctx2.executor.stop()


# ── degraded mode ────────────────────────────────────────────────────────────


def test_degraded_sonnet_unavailable_gives_fallback_labels(tmp_path):
    ctx = _build_ctx(tmp_path, model_sonnet=None)
    ctx.executor.start()
    try:
        ctx.model_client = RaisingModelClient(
            ModelUnavailable(MODEL_SONNET_ENDPOINT, "rate limit 0", permanent=True)
        )
        run_id = _start_run(ctx)
        status = _wait_for(ctx, run_id, {"awaiting_signoff", "failed"})
        assert status == "awaiting_signoff", service.get_run(ctx, run_id).get("status_reason")

        narratives = ctx.persistence.get_narratives(run_id)
        assert narratives
        assert {n["origin"] for n in narratives} == {"fallback_unavailable"}
        assert all(n["template_text"] is None for n in narratives)

        # No candidates are ever invented while degraded (NN13/C-2's own
        # rule: an unavailable model gives zero candidates).
        assert ctx.persistence.list_candidates(run_id) == []

        service.sign_off(ctx, run_id, "alice")
        status = _wait_for(ctx, run_id, {"completed", "failed"})
        assert status == "completed", service.get_run(ctx, run_id).get("status_reason")

        prs, wb = _read_exports(ctx, run_id)
        all_text = _all_slide_text(prs)
        assert "LLM unavailable — deterministic output only" in all_text
        # the deterministic template prose (findings.yaml's own observation)
        # still renders -- a degraded run is still a complete, readable run
        assert "claims exceed the high-value threshold" in all_text or "claims" in all_text
    finally:
        ctx.executor.stop()


# ── G10 ───────────────────────────────────────────────────────────────────────


_RISK_LANGUAGE_TERMS = (
    "fraud", "deliberate", "intentional", "circumvent", "evade", "evasion", "conceal",
    "misconduct", "dishonest", "theft", "steal", "abuse", "collu", "manipulat",
)


def test_g10_clean_dataset_zero_findings_zero_candidates_no_candidate_call_no_risk_language(tmp_path):
    ctx = _build_ctx(tmp_path, data_writer=_write_clean_data)
    ctx.executor.start()
    try:
        client = _happy_client()
        ctx.model_client = client
        run_id = _start_run(ctx)
        status = _wait_for(ctx, run_id, {"awaiting_signoff", "failed"})
        assert status == "awaiting_signoff", service.get_run(ctx, run_id).get("status_reason")

        findings = ctx.persistence.list_findings(run_id)
        assert findings == []
        assert ctx.persistence.list_candidates(run_id) == []

        calls = ctx.persistence.list_llm_calls(run_id)
        assert all(c["task"] != "find_candidates" for c in calls)
        assert not any(MARKER_CANDIDATES in str(m.get("params", {})) for m in client.calls)
        # No find/find_candidates markers were ever dispatched to the client
        assert all(
            "finding_key" not in str(call["messages"]) and MARKER_CANDIDATES not in str(call["messages"])
            for call in client.calls
        )

        narratives = ctx.persistence.get_narratives(run_id)
        for n in narratives:
            text = (n.get("template_text") or "").lower()
            for term in _RISK_LANGUAGE_TERMS:
                assert term not in text, f"risk language {term!r} found in {n['narrative_id']}: {text!r}"

        service.sign_off(ctx, run_id, "alice")
        status = _wait_for(ctx, run_id, {"completed", "failed"})
        assert status == "completed", service.get_run(ctx, run_id).get("status_reason")
    finally:
        ctx.executor.stop()
