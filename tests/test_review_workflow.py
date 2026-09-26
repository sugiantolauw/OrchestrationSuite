"""P7 review workflow end-to-end (docs/specs/P7_mapping_authoring_design.md
§3.3-3.4, §3.10): preparer -> reviewer -> approver, review notes, return to
preparer, racing "Mark as reviewed", a replayed step, and the preparation-
only edit/candidate-decision gate (D-P7-10)."""

from __future__ import annotations

import dataclasses
import threading

import pytest

from orchestrator import service
from orchestrator.errors import (
    CandidateSeverityRequired,
    ReviewActionRefused,
    RunNotAwaitingSignoff,
    StaleStateError,
)
from tests.n9_test_support import harness_at_awaiting_signoff


def _config_ctx(ctx_app, tmp_path, roles: dict[str, list[str]], *, sod_mode="enforced"):
    path = tmp_path / "review_roles.yaml"
    lines = "\n".join(f"{email}: [{', '.join(rs)}]" for email, rs in roles.items())
    path.write_text(lines + "\n")
    settings = dataclasses.replace(
        ctx_app.settings,
        review_role_source="config",
        review_role_assignments_path=str(path),
        review_sod_mode=sod_mode,
        review_preparer_groups=("audit-preparers",),
        review_reviewer_groups=("audit-reviewers",),
        review_approver_groups=("audit-approvers",),
    )
    return dataclasses.replace(ctx_app, settings=settings)


_THREE_ROLES = {
    "alice": ["preparer"],
    "bob": ["reviewer"],
    "carol": ["approver"],
}


def test_happy_path_prepare_review_approve(local_persistence, tmp_path, clock):
    h, ctx_app, state = harness_at_awaiting_signoff(local_persistence, tmp_path, clock)
    run_id = state.run_id
    ctx = _config_ctx(ctx_app, tmp_path, _THREE_ROLES)

    prepared = service.prepare_findings(ctx, run_id, "alice")
    assert prepared.review["stage"] == "review"
    assert prepared.review["prepared"]["actor"] == "alice"
    assert all(f["review_state"] == "prepared" for f in h.persistence.list_findings(run_id))

    reviewed = service.mark_reviewed(ctx, run_id, "bob")
    assert reviewed.review["stage"] == "approval"
    assert reviewed.review["reviewed"]["actor"] == "bob"
    assert all(f["review_state"] == "reviewed" for f in h.persistence.list_findings(run_id))

    signed = service.sign_off(ctx, run_id, "carol")
    assert signed.status == "completed" or signed.status == "queued"
    assert signed.signoff["sod_mode"] == "enforced"
    assert signed.signoff["prepared_by"] == "alice"
    assert signed.signoff["reviewed_by"] == "bob"
    assert signed.signoff["approver"] == "carol"
    assert signed.signoff["self_approved"] is False
    assert all(f["review_state"] == "approved" for f in h.persistence.list_findings(run_id))

    steps = h.persistence.list_review_steps(run_id)
    assert [s["action"] for s in steps] == ["prepared", "reviewed", "approved"]
    assert steps[0]["role_source"] == "config"
    assert steps[0]["sod_mode"] == "enforced"

    runs_row = next(r for r in h.persistence.list_runs() if r["run_id"] == run_id)
    assert runs_row["prepared_by"] == "alice"
    assert runs_row["reviewed_by"] == "bob"
    assert runs_row["approved_by"] == "carol"


def test_wrong_role_refused_with_group_names_in_message(local_persistence, tmp_path, clock):
    h, ctx_app, state = harness_at_awaiting_signoff(local_persistence, tmp_path, clock)
    run_id = state.run_id
    ctx = _config_ctx(ctx_app, tmp_path, _THREE_ROLES)
    service.prepare_findings(ctx, run_id, "alice")

    with pytest.raises(ReviewActionRefused) as exc:
        service.mark_reviewed(ctx, run_id, "alice")  # alice is only a preparer
    assert "not in a preparer/reviewer/approver group" in exc.value.reason
    assert "audit-reviewers" in exc.value.reason


def test_refused_trace_events_keyed_by_actor_not_just_state_version(local_persistence, tmp_path, clock):
    # BUG-P2B-1 (independent review, RUN-69A978937B2E): a refusal never
    # advances RunState, so two different actors refused back-to-back are
    # refused at the SAME state_version -- the trace event id must still
    # differ between them, or the second silently no-ops into the first's
    # row (append_trace_event dedups by event_id).
    h, ctx_app, state = harness_at_awaiting_signoff(local_persistence, tmp_path, clock)
    run_id = state.run_id
    ctx = _config_ctx(ctx_app, tmp_path, {"alice": ["preparer"], "bob": ["preparer"]})
    service.prepare_findings(ctx, run_id, "alice")

    with pytest.raises(ReviewActionRefused):
        service.mark_reviewed(ctx, run_id, "alice")  # alice is only a preparer
    with pytest.raises(ReviewActionRefused):
        service.mark_reviewed(ctx, run_id, "bob")  # bob is only a preparer too, same state_version

    refusals = [e for e in h.persistence.list_trace_events(run_id) if e["event_type"] == "review_action_refused"]
    assert sorted(e["actor"] for e in refusals) == ["alice", "bob"]

    # An identical retry by the same actor (same action, same state_version)
    # stays idempotent -- no third row.
    with pytest.raises(ReviewActionRefused):
        service.mark_reviewed(ctx, run_id, "alice")
    refusals_after_retry = [
        e for e in h.persistence.list_trace_events(run_id) if e["event_type"] == "review_action_refused"
    ]
    assert len(refusals_after_retry) == 2


def test_sod_enforced_refuses_second_role_for_same_actor(local_persistence, tmp_path, clock):
    h, ctx_app, state = harness_at_awaiting_signoff(local_persistence, tmp_path, clock)
    run_id = state.run_id
    ctx = _config_ctx(ctx_app, tmp_path, {"alice": ["preparer", "reviewer"]})

    service.prepare_findings(ctx, run_id, "alice")
    with pytest.raises(ReviewActionRefused) as exc:
        service.mark_reviewed(ctx, run_id, "alice")
    assert "already acted on this run as preparer" in exc.value.reason


def test_sod_labelled_waives_second_role_for_same_actor(local_persistence, tmp_path, clock):
    h, ctx_app, state = harness_at_awaiting_signoff(local_persistence, tmp_path, clock)
    run_id = state.run_id
    ctx = _config_ctx(ctx_app, tmp_path, {"alice": ["preparer", "reviewer", "approver"]}, sod_mode="labelled")

    service.prepare_findings(ctx, run_id, "alice")
    service.mark_reviewed(ctx, run_id, "alice")
    signed = service.sign_off(ctx, run_id, "alice")

    assert signed.signoff["self_approved"] is True
    assert signed.signoff["sod_mode"] == "labelled"


def test_open_note_blocks_review_and_signoff(local_persistence, tmp_path, clock):
    h, ctx_app, state = harness_at_awaiting_signoff(local_persistence, tmp_path, clock)
    run_id = state.run_id
    ctx = _config_ctx(ctx_app, tmp_path, _THREE_ROLES)

    service.prepare_findings(ctx, run_id, "alice")
    note = service.raise_review_note(ctx, run_id, "bob", "justify excluding this", None)
    assert note["state"] == "open"

    with pytest.raises(ReviewActionRefused) as exc:
        service.mark_reviewed(ctx, run_id, "bob")
    assert "1 open" in exc.value.reason

    service.respond_to_review_note(ctx, run_id, note["note_id"], "alice", "because X")
    with pytest.raises(ReviewActionRefused):
        service.mark_reviewed(ctx, run_id, "bob")

    service.clear_review_note(ctx, run_id, note["note_id"], "bob")
    reviewed = service.mark_reviewed(ctx, run_id, "bob")
    assert reviewed.review["stage"] == "approval"


def test_return_resets_findings_to_draft_and_keeps_open_notes(local_persistence, tmp_path, clock):
    h, ctx_app, state = harness_at_awaiting_signoff(local_persistence, tmp_path, clock)
    run_id = state.run_id
    ctx = _config_ctx(ctx_app, tmp_path, _THREE_ROLES)

    service.prepare_findings(ctx, run_id, "alice")
    note = service.raise_review_note(ctx, run_id, "bob", "justify this", None)

    returned = service.return_to_preparer(ctx, run_id, "bob", "please justify T1")

    assert returned.review["stage"] == "preparation"
    assert returned.review["returns"][-1]["reason"] == "please justify T1"
    assert all(f["review_state"] == "draft" for f in h.persistence.list_findings(run_id))
    # open notes stay open across a return
    still_open = [n for n in h.persistence.list_review_notes(run_id) if n["note_id"] == note["note_id"]]
    assert still_open[0]["state"] == "open"


def test_racing_mark_reviewed_exactly_one_winner(local_persistence, tmp_path, clock):
    h, ctx_app, state = harness_at_awaiting_signoff(local_persistence, tmp_path, clock)
    run_id = state.run_id
    ctx = _config_ctx(ctx_app, tmp_path, {"alice": ["preparer"], "bob": ["reviewer"]})
    service.prepare_findings(ctx, run_id, "alice")

    barrier = threading.Barrier(2)
    results = []

    def attempt():
        barrier.wait()
        try:
            r = service.mark_reviewed(ctx, run_id, "bob")
            results.append(("ok", r.review["stage"]))
        except (StaleStateError, ReviewActionRefused):
            # Both are the accepted "lost the race" outcomes (same shape as
            # tests/test_failure_injection.py's resume race): a CAS
            # rejection (the loser read stage='review' but lost save_state's
            # CAS), or a stage-mismatch refusal (the loser's load_state ran
            # late enough to see the winner's already-advanced stage).
            results.append(("lost", None))

    threads = [threading.Thread(target=attempt) for _ in range(2)]
    for t in threads:
        t.start()
    for t in threads:
        t.join()

    oks = [r for r in results if r[0] == "ok"]
    losers = [r for r in results if r[0] == "lost"]
    assert len(oks) == 1
    assert len(losers) == 1

    reviewed_steps = [s for s in h.persistence.list_review_steps(run_id) if s["action"] == "reviewed"]
    assert len(reviewed_steps) == 1


def test_a_replayed_prepare_step_does_not_duplicate(local_persistence, tmp_path, clock):
    h, ctx_app, state = harness_at_awaiting_signoff(local_persistence, tmp_path, clock)
    run_id = state.run_id
    ctx = _config_ctx(ctx_app, tmp_path, _THREE_ROLES)

    service.prepare_findings(ctx, run_id, "alice")
    with pytest.raises(ReviewActionRefused):
        service.prepare_findings(ctx, run_id, "alice")

    prepared_steps = [s for s in h.persistence.list_review_steps(run_id) if s["action"] == "prepared"]
    assert len(prepared_steps) == 1


def test_edit_after_preparation_refused(local_persistence, tmp_path, clock):
    from tests.test_g14_narrative_edits import _t1_observation_row

    h, ctx_app, state = harness_at_awaiting_signoff(local_persistence, tmp_path, clock)
    run_id = state.run_id
    ctx = _config_ctx(ctx_app, tmp_path, _THREE_ROLES)

    row = _t1_observation_row(h.persistence, run_id)
    service.prepare_findings(ctx, run_id, "alice")

    with pytest.raises(ReviewActionRefused) as exc:
        service.edit_narrative(ctx, run_id, row["narrative_id"], "edited text", actor="alice")
    assert "ask the reviewer to return the run" in exc.value.reason


def test_decide_candidate_after_preparation_refused(local_persistence, tmp_path, clock):
    h, ctx_app, state = harness_at_awaiting_signoff(local_persistence, tmp_path, clock)
    run_id = state.run_id
    ctx = _config_ctx(ctx_app, tmp_path, _THREE_ROLES)
    h.persistence.write_candidates(
        run_id,
        [{
            "candidate_id": "C1", "engagement_id": "ENG-DEFAULT", "skill_id": "SKILL-MINI",
            "generation": 0, "rule_id": "SKILL-MINI.ai.abc123", "title": "A candidate finding",
            "metrics_cited": ["hv_count"], "producing_test_ids": ["T1"], "proposed_severity": "Medium",
            "severity_reason": "reason", "rationale": "rationale", "monetary_basis": "none",
            "monetary_basis_note": "no monetary basis", "exposure_amount": None,
            "headline_eligible": False, "headline_ineligible_reason": "monetary_basis is none",
            "candidate_status": "candidate", "call_id": "CALL-C1",
        }],
        now=clock(),
    )
    service.decide_candidate(ctx, run_id, "C1", decision="accepted", reason=None, decided_severity="Low", actor="alice")
    service.prepare_findings(ctx, run_id, "alice")

    with pytest.raises(ReviewActionRefused) as exc:
        service.decide_candidate(ctx, run_id, "C1", decision="rejected", reason="x", decided_severity=None, actor="alice")
    assert "ask the reviewer to return the run" in exc.value.reason


def test_theme_generation_recorded_at_prepare(local_persistence, tmp_path, clock):
    h, ctx_app, state = harness_at_awaiting_signoff(local_persistence, tmp_path, clock)
    run_id = state.run_id
    ctx = _config_ctx(ctx_app, tmp_path, _THREE_ROLES)

    current = h.persistence.load_state(run_id)
    bumped = dataclasses.replace(current, options={**current.options, "narration_generation": 2})
    h.persistence.save_state(bumped)

    service.prepare_findings(ctx, run_id, "alice")

    [step] = [s for s in h.persistence.list_review_steps(run_id) if s["action"] == "prepared"]
    assert step["theme_generation"] == 2
    assert step["narration_generation"] == 2


def test_prepare_refused_with_undecided_candidates(local_persistence, tmp_path, clock):
    h, ctx_app, state = harness_at_awaiting_signoff(local_persistence, tmp_path, clock)
    run_id = state.run_id
    ctx = _config_ctx(ctx_app, tmp_path, _THREE_ROLES)
    h.persistence.write_candidates(
        run_id,
        [{
            "candidate_id": "C2", "engagement_id": "ENG-DEFAULT", "skill_id": "SKILL-MINI",
            "generation": 0, "rule_id": "SKILL-MINI.ai.def456", "title": "Undecided",
            "metrics_cited": ["hv_count"], "producing_test_ids": ["T1"], "proposed_severity": "Medium",
            "severity_reason": "reason", "rationale": "rationale", "monetary_basis": "none",
            "monetary_basis_note": "no monetary basis", "exposure_amount": None,
            "headline_eligible": False, "headline_ineligible_reason": "monetary_basis is none",
            "candidate_status": "candidate", "call_id": "CALL-C2",
        }],
        now=clock(),
    )
    from orchestrator.errors import CandidatesUndecided

    with pytest.raises(CandidatesUndecided):
        service.prepare_findings(ctx, run_id, "alice")


def test_actions_refused_when_run_not_awaiting_signoff(local_persistence, tmp_path, clock):
    h, ctx_app, state = harness_at_awaiting_signoff(local_persistence, tmp_path, clock)
    run_id = state.run_id
    ctx = _config_ctx(ctx_app, tmp_path, _THREE_ROLES)
    service.prepare_findings(ctx, run_id, "alice")
    service.mark_reviewed(ctx, run_id, "bob")
    service.sign_off(ctx, run_id, "carol")

    with pytest.raises(RunNotAwaitingSignoff):
        service.prepare_findings(ctx, run_id, "alice")


# ── Failure injection (§3.10, CLAUDE.md §9C: CAS + idempotent MERGE) ────────
#
# orchestrator.runs.prepare/mark_reviewed/return_to_preparer write in this
# order: (1) the CAS `save_state` (the run's `review` stage), THEN (2) the
# finding review_state advance, THEN (3) the `review_steps` evidence row --
# CAS first, deliberately, so two racing actors can never both advance the
# same findings (test_racing_mark_reviewed_exactly_one_winner above is the
# proof: finding-advancement AFTER the CAS lock is what makes a loser fail
# with StaleStateError instead of a raw InvalidReviewStateTransition).
# A crash between (1) and (3) therefore leaves `state.review` already
# advanced with the matching `review_steps` row still missing -- this test
# proves that gap is recoverable, not corrupting: `append_review_step`'s
# deterministic step_id (keyed on the PRE-transition state_version) lets a
# repair pass write the same row a second time as a safe no-op, and a raw
# retry of the same action is cleanly refused (wrong stage) rather than
# double-applying.


def test_crash_after_save_state_before_review_step_is_recoverable(local_persistence, tmp_path, clock):
    from orchestrator import runs as runs_module

    h, ctx_app, state = harness_at_awaiting_signoff(local_persistence, tmp_path, clock)
    run_id = state.run_id
    ctx = _config_ctx(ctx_app, tmp_path, _THREE_ROLES)

    real_append = h.persistence.append_review_step
    calls = {"n": 0}

    def crashing_append(row):
        calls["n"] += 1
        if calls["n"] == 1:
            raise RuntimeError("simulated crash after save_state, before append_review_step")
        return real_append(row)

    h.persistence.append_review_step = crashing_append
    with pytest.raises(RuntimeError):
        service.prepare_findings(ctx, run_id, "alice")
    h.persistence.append_review_step = real_append

    # (1) already committed: the run's review stage advanced despite the crash.
    crashed_state = h.persistence.load_state(run_id)
    assert crashed_state.review["stage"] == "review"
    assert crashed_state.review["prepared"]["actor"] == "alice"
    # (3) is missing.
    assert h.persistence.list_review_steps(run_id) == []

    # A raw retry of the same action is cleanly refused (wrong stage now),
    # never a duplicate or a corrupted finding-state walk.
    with pytest.raises(ReviewActionRefused):
        service.prepare_findings(ctx, run_id, "alice")

    # A repair pass reconstructs and writes the SAME row the crash lost --
    # the deterministic step_id (keyed on the pre-transition state_version,
    # `state.state_version` as it was BEFORE prepare()'s own save_state)
    # makes this idempotent even if the repair itself is retried.
    missing_row = {
        "step_id": runs_module._step_id(run_id, "prepared", state.state_version),
        "run_id": run_id, "engagement_id": state.engagement_id, "action": "prepared",
        "actor": "alice", "role": "preparer", "matched_group": "audit-preparers",
        "role_source": "config", "sod_mode": "enforced",
        "reason": None, "theme_generation": 0, "narration_generation": 0,
        "state_version": state.state_version, "at": clock(),
    }
    h.persistence.append_review_step(missing_row)
    h.persistence.append_review_step(missing_row)  # retried repair -- still idempotent

    steps = h.persistence.list_review_steps(run_id)
    assert len(steps) == 1
    assert steps[0]["action"] == "prepared"
    assert steps[0]["actor"] == "alice"


# ── G13 export fidelity (§3.7) ───────────────────────────────────────────────


def test_xlsx_cover_and_review_notes_sheet_match_the_run(local_persistence, tmp_path, clock):
    import io

    import openpyxl

    from orchestrator.nodes.fieldwork import export

    h, ctx_app, state = harness_at_awaiting_signoff(local_persistence, tmp_path, clock)
    run_id = state.run_id
    ctx = _config_ctx(ctx_app, tmp_path, _THREE_ROLES)

    service.prepare_findings(ctx, run_id, "alice")
    note = service.raise_review_note(ctx, run_id, "bob", "please justify T1", None)
    service.respond_to_review_note(ctx, run_id, note["note_id"], "alice", "because X")
    service.clear_review_note(ctx, run_id, note["note_id"], "bob")
    reviewed = service.mark_reviewed(ctx, run_id, "bob")
    signed = service.sign_off(ctx, run_id, "carol")

    exported_state = export(h.ctx, signed)
    xlsx_meta = exported_state.exports["xlsx"]
    content = h.ctx.export_storage.read(xlsx_meta["path"])
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)

    cover = wb["Cover"]
    cover_rows = {row[0].value: row[1].value for row in cover.iter_rows(min_row=2, max_col=2) if row[0].value}
    assert cover_rows["signed_off_by"] == "carol"
    assert cover_rows["prepared_by"] == "alice"
    assert cover_rows["reviewed_by"] == "bob"
    assert cover_rows["sod_mode"] == "enforced"
    assert cover_rows["role_source"] == "config"
    assert "prepared_at" in cover_rows
    assert "reviewed_at" in cover_rows
    assert "signoff_note" not in cover_rows  # SoD not waived here

    notes_sheet = wb["Review notes"]
    header = [c.value for c in notes_sheet[1]]
    rows = [
        dict(zip(header, [c.value for c in r]))
        for r in notes_sheet.iter_rows(min_row=2)
        if r[header.index("raised_by")].value
    ]
    assert len(rows) == 1
    row = rows[0]
    assert row["body"] == "please justify T1"
    assert row["target"] == "Whole run"
    assert row["raised_by"] == "bob"
    assert row["raised_role"] == "reviewer"
    assert row["state"] == "cleared"
    assert row["response"] == "because X"
    assert row["responded_by"] == "alice"
    assert row["cleared_by"] == "bob"
    assert row["cleared_role"] == "reviewer"


def test_pptx_signoff_line_and_trace_message(local_persistence, tmp_path, clock):
    from orchestrator.nodes.fieldwork import export
    from orchestrator.pptx_export import _signoff_line

    h, ctx_app, state = harness_at_awaiting_signoff(local_persistence, tmp_path, clock)
    run_id = state.run_id
    ctx = _config_ctx(ctx_app, tmp_path, _THREE_ROLES)

    service.prepare_findings(ctx, run_id, "alice")
    service.mark_reviewed(ctx, run_id, "bob")
    signed = service.sign_off(ctx, run_id, "carol")

    line = _signoff_line(signed)
    assert line.startswith("Prepared by alice · Reviewed by bob · Signed off by carol on")
    assert "Self-approved" not in line

    exported_state = export(h.ctx, signed)
    assert exported_state.exports["pptx"]["path"]

    events = [e for e in h.persistence.list_trace_events(run_id) if e["event_type"] == "signed_off"]
    assert len(events) == 1
    assert events[0]["message"] == "Findings signed off by carol (prepared by alice, reviewed by bob)"


def test_legacy_direct_sign_off_unaffected(local_persistence, tmp_path, clock):
    """A caller that never enters the P7 workflow at all (state.review stays
    None) keeps getting the exact pre-P7 self-sign-off behaviour -- the
    dual-path design D-P7-9 relies on."""
    h, ctx_app, state = harness_at_awaiting_signoff(local_persistence, tmp_path, clock)
    run_id = state.run_id

    signed = service.sign_off(ctx_app, run_id, state.run_owner)
    assert "sod_mode" not in signed.signoff
    assert signed.signoff["self_approved"] is True
    assert all(f["review_state"] == "approved" for f in h.persistence.list_findings(run_id))
