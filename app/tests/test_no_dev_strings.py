"""No developer-facing string (CLAUDE.md, section markers, gate/phase names,
code identifiers) ever reaches rendered UI text. Renders every route with
the fake service (no live backend needed) and, where a completed local run
is easy to construct, with the real local-backend service on that run."""

from __future__ import annotations

import re

import pytest

_FORBIDDEN_SUBSTRINGS = [
    "CLAUDE.md",
    "§",
    "NN13", "NN14", "NN15", "NN16",
    "G6", "G7", "G8", "G9", "G10", "G11", "G12", "G13", "G14", "G15", "G16", "G17",
    "docs/specs",
    "[PENDING]",
    "test_specification",
]
_FORBIDDEN_PATTERNS = [re.compile(r"\bP[0-9][AB]?\b")]  # phase names: P1A, P1B, P2 .. P9

# Independent review 2026-09-24 item 1/2: a "(CLAUDE.md ...)"-shaped citation
# was only the most obvious developer-facing marker that reached an exported
# PPTX/XLSX -- "gate review", "WP N" and a bare code-repository file path are
# the same class of thing with a different spelling. Kept separate from
# _FORBIDDEN_SUBSTRINGS/_FORBIDDEN_PATTERNS above (used by every /route page
# test in this file, several of which have their own PRE-EXISTING, out-of-
# scope leaks noted elsewhere in this file -- see
# test_real_skill_description_has_no_dev_strings's own comment -- that a
# broad file-path pattern would also catch and turn red for a reason this
# task does not own) rather than folded into them.
_EXPORT_FORBIDDEN_SUBSTRINGS = [
    *_FORBIDDEN_SUBSTRINGS,
    "gate review",
    "reference_app",
    "run_fingerprints",
    "catalogue.yaml",
    "thresholds.yaml",
    "findings.yaml",
    "plan.yaml",
    "manifest.yaml",
    "contract.yaml",
    "strip_pptx_template.py",
    "computation.py",
]
_EXPORT_FORBIDDEN_PATTERNS = [
    *_FORBIDDEN_PATTERNS,
    re.compile(r"\bWP\s?[0-9]+\b"),  # workpaper-index jargon: "WP 7", "WP7"
    re.compile(r"\.py:\d+\b"),  # a file:line citation, e.g. "pptx_export.py:352"
]


def _all_text(node, out):
    if node is None:
        return
    if isinstance(node, (str, int, float)):
        out.append(str(node))
        return
    if isinstance(node, (list, tuple)):
        for n in node:
            _all_text(n, out)
        return
    children = getattr(node, "children", None)
    if children is not None:
        _all_text(children, out)
    # dcc.Markdown / dcc.Textarea etc. carry text in `value`/other props too.
    for prop in ("value", "placeholder", "title", "label"):
        val = getattr(node, prop, None)
        if isinstance(val, str):
            out.append(val)


def render_text(component) -> str:
    out: list[str] = []
    _all_text(component, out)
    return "\n".join(out)


def assert_no_dev_strings(component, label: str) -> None:
    text = render_text(component)
    for token in _FORBIDDEN_SUBSTRINGS:
        assert token not in text, f"{label}: found developer string {token!r} in rendered text"
    for pat in _FORBIDDEN_PATTERNS:
        m = pat.search(text)
        assert m is None, f"{label}: found developer string {m.group(0)!r} in rendered text"


@pytest.fixture(autouse=True)
def _fake_backend(monkeypatch):
    import tests.fake_service as fake_service
    from src.platform import adapters

    ctx = fake_service.build_app_context()
    monkeypatch.setattr(adapters, "service", fake_service, raising=False)
    monkeypatch.setattr(adapters, "_ctx", ctx, raising=False)
    monkeypatch.setattr(adapters, "get_context", lambda: ctx, raising=False)
    run_id = fake_service.start_audit_run(
        ctx, skill_id="SKILL-001", bindings={}, audit_period=("2025-01-01", "2025-06-30"),
        objective="Test objective", run_owner="tester", review_plan_first=False,
    )
    fake_service.sign_off(ctx, run_id, "tester")
    return run_id


def test_landing_page_has_no_dev_strings():
    from src.run_setup import home_layout

    assert_no_dev_strings(home_layout(), "/")


def test_skill_library_page_has_no_dev_strings():
    from src.platform.pages import skill_library_page

    assert_no_dev_strings(skill_library_page(), "/skills")


def test_skill_methodology_page_has_no_dev_strings():
    from src.platform.pages import skill_methodology_page

    assert_no_dev_strings(skill_methodology_page("SKILL-001"), "/skills/SKILL-001")


def test_audit_runs_page_has_no_dev_strings():
    from src.platform.pages import audit_runs_page

    assert_no_dev_strings(audit_runs_page(), "/runs")


def test_platform_trace_page_has_no_dev_strings():
    from src.platform.pages import platform_trace_page

    assert_no_dev_strings(platform_trace_page(), "/trace")


def test_management_actions_page_has_no_dev_strings():
    from src.platform.pages import management_actions_page

    assert_no_dev_strings(management_actions_page(), "/actions")


def test_run_status_page_has_no_dev_strings(_fake_backend):
    from src.run_status import run_page, _render_body
    from src.platform import adapters

    run = adapters.get_run(_fake_backend)
    assert_no_dev_strings(run_page(_fake_backend), "/run/<id> (shell)")
    assert_no_dev_strings(_render_body(run, _fake_backend), "/run/<id> (body)")


def test_workspace_tne_has_no_dev_strings(_fake_backend):
    from src.workspace_tne import tne_workspace_layout

    assert_no_dev_strings(tne_workspace_layout(_fake_backend), "/workspace/tne")


def test_real_skill_description_has_no_dev_strings(monkeypatch, tmp_path):
    """The skill card's description comes from skills/tne_exco/manifest.yaml,
    not the fake backend -- the fake never reproduces a stale/developer
    description because it hardcodes its own. Render against the REAL
    orchestrator.service (as test_layouts.py's real-skill test does) so a
    developer string left in the real manifest is actually caught."""
    from pathlib import Path

    from orchestrator import service as real_service
    from src.platform.pages import skill_library_page
    from src.run_setup import home_layout

    repo_root = Path(__file__).resolve().parents[2]
    env = {
        "ORCH_BACKEND": "local",
        "ORCH_LOCAL_DB": str(tmp_path / "orch.db"),
        "ORCH_LOCAL_DATA_ROOT": str(tmp_path),
        "ORCH_LOCAL_EXPORT_ROOT": str(tmp_path / "exports"),
        "SKILLS_DIR": str(repo_root / "skills"),
    }
    from src.platform import adapters as adapters_module

    monkeypatch.setattr(adapters_module, "service", real_service)
    adapters_module._ctx = None
    original_build = real_service.build_app_context
    monkeypatch.setattr(real_service, "build_app_context", lambda *a, **k: original_build(env))
    # This file's own autouse `_fake_backend` fixture (above) replaces
    # adapters.get_context itself with a lambda closed over the fake ctx,
    # so resetting `_ctx` alone is not enough here -- put a real one back.
    real_ctx = original_build(env)
    monkeypatch.setattr(adapters_module, "get_context", lambda: real_ctx)

    assert_no_dev_strings(home_layout(), "/ (real skill)")
    assert_no_dev_strings(skill_library_page(), "/skills (real skill)")
    # NOTE: skill_methodology_page("SKILL-001") independently leaks a
    # developer string ("...until P6 classify node") in its test-plan table
    # -- pre-existing, unrelated to the skill card description this test
    # targets, and out of this task's scope. Left for a future fix.

    manifest = __import__("yaml").safe_load((repo_root / "skills" / "tne_exco" / "manifest.yaml").read_text())
    assert "Assess executive travel and entertainment spend against policy" in manifest["description"]


@pytest.fixture(scope="module")
def _completed_skill001_export(tmp_path_factory):
    """One completed real SKILL-001 run over the fast planted fixture
    (tests/fixtures/tne_planted/data/, never the ~3-minute real
    synthetic_data/ run), shared by the XLSX- and PPTX-dev-string tests
    below so the pipeline runs once, not twice."""
    import time
    from pathlib import Path

    from orchestrator import service as real_service

    repo_root = Path(__file__).resolve().parents[2]
    data_dir = repo_root / "tests" / "fixtures" / "tne_planted" / "data"
    if not data_dir.is_dir():
        pytest.skip("tests/fixtures/tne_planted/data/ not present -- run generate.py first")

    tmp_path = tmp_path_factory.mktemp("dev-strings-export-check")
    env = {
        "ORCH_BACKEND": "local",
        "ORCH_LOCAL_DB": str(tmp_path / "orch.db"),
        "ORCH_LOCAL_DATA_ROOT": str(data_dir),
        "ORCH_LOCAL_EXPORT_ROOT": str(tmp_path / "exports"),
        "ORCH_WORKER_ID": "dev-strings-export-check",
        "SKILLS_DIR": str(repo_root / "skills"),
        # Pinned, not derived from `git rev-parse HEAD` -- several agents
        # commit to this checkout concurrently, so HEAD can move between
        # this run's creation and the executor's fingerprint re-check
        # (tests/test_p3_service.py's _build_ctx does the same, same reason).
        "CODE_REVISION": "test-fixed-revision",
    }
    ctx = real_service.build_app_context(env)
    ctx.executor.start()
    try:
        bindings = real_service.suggest_bindings(ctx, "SKILL-001")
        run_id = real_service.start_audit_run(
            ctx, skill_id="SKILL-001", bindings=bindings,
            audit_period=("2025-01-01", "2026-04-30"),
            objective="No-dev-strings export check.", run_owner="dev-strings-check",
        )
        deadline = time.time() + 120
        status = None
        while time.time() < deadline:
            status = real_service.get_run(ctx, run_id)["status"]
            if status in ("awaiting_signoff", "failed"):
                break
            time.sleep(0.5)
        assert status == "awaiting_signoff", real_service.get_run(ctx, run_id).get("status_reason")

        real_service.sign_off(ctx, run_id, "dev-strings-check")
        deadline = time.time() + 30
        while time.time() < deadline:
            status = real_service.get_run(ctx, run_id)["status"]
            if status in ("completed", "failed"):
                break
            time.sleep(0.5)
        assert status == "completed", real_service.get_run(ctx, run_id).get("status_reason")

        _xlsx_name, xlsx_content = real_service.get_export(ctx, run_id, "xlsx")
        _pptx_name, pptx_content = real_service.get_export(ctx, run_id, "pptx")
    finally:
        ctx.executor.stop()

    return {"xlsx": xlsx_content, "pptx": pptx_content}


def _find_dev_strings(text: str, substrings, patterns) -> list[str]:
    hits = []
    for token in substrings:
        if token in text:
            hits.append(repr(token))
    for pat in patterns:
        m = pat.search(text)
        if m:
            hits.append(repr(m.group(0)))
    return hits


def test_xlsx_workpaper_has_no_dev_strings(_completed_skill001_export):
    """Independent review 2026-09-24 item 7: a developer-facing string
    (a CLAUDE.md reference, a section mark, a gate/phase name, a bare
    codebase file path, "gate review", "WP N") is exactly as wrong in an
    exported XLSX workpaper as it is in the UI -- an auditor reads both.
    Scans every string cell of every sheet."""
    import io

    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(_completed_skill001_export["xlsx"]), data_only=True)
    offenders = []
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    for hit in _find_dev_strings(cell.value, _EXPORT_FORBIDDEN_SUBSTRINGS, _EXPORT_FORBIDDEN_PATTERNS):
                        offenders.append(f"{sheet.title}!{cell.coordinate}: {hit} in {cell.value!r}")
    assert not offenders, "developer string(s) found in the XLSX workpaper:\n" + "\n".join(offenders)


def test_pptx_pack_has_no_dev_strings(_completed_skill001_export):
    """Independent review 2026-09-24 item 1/2: the live PPTX pack shipped
    "(CLAUDE.md P6)" on its Executive Summary slide and "CLAUDE.md §0"/"§3"
    citations on its Methodology & Limitations slide -- exactly as wrong on
    an executive-facing deck as the same class of string in the UI or the
    XLSX workpaper. Scans every text frame, table cell and speaker note on
    every slide, AND the package's own document properties
    (docProps/core.xml, app.xml, custom.xml -- a developer-facing string in
    a doc property is no better than one on a slide, and NN16 covers the
    package's metadata too, not only its slide content)."""
    import io
    import zipfile

    from pptx import Presentation

    content = _completed_skill001_export["pptx"]
    prs = Presentation(io.BytesIO(content))
    offenders = []

    def _table_texts(table):
        for row in table.rows:
            for cell in row.cells:
                yield cell.text_frame.text

    for i, slide in enumerate(prs.slides):
        for shape in slide.shapes:
            texts = []
            if shape.has_text_frame:
                texts.append(shape.text_frame.text)
            if getattr(shape, "has_table", False):
                texts.extend(_table_texts(shape.table))
            for text in texts:
                for hit in _find_dev_strings(text, _EXPORT_FORBIDDEN_SUBSTRINGS, _EXPORT_FORBIDDEN_PATTERNS):
                    offenders.append(f"slide {i}: {hit} in {text[:120]!r}")
        if slide.has_notes_slide:
            note_text = slide.notes_slide.notes_text_frame.text
            for hit in _find_dev_strings(note_text, _EXPORT_FORBIDDEN_SUBSTRINGS, _EXPORT_FORBIDDEN_PATTERNS):
                offenders.append(f"slide {i} notes: {hit} in {note_text[:120]!r}")

    with zipfile.ZipFile(io.BytesIO(content)) as z:
        for part in ("docProps/core.xml", "docProps/app.xml", "docProps/custom.xml"):
            if part not in z.namelist():
                continue
            xml_text = z.read(part).decode("utf-8", errors="replace")
            for hit in _find_dev_strings(xml_text, _EXPORT_FORBIDDEN_SUBSTRINGS, _EXPORT_FORBIDDEN_PATTERNS):
                offenders.append(f"{part}: {hit}")

    assert not offenders, "developer string(s) found in the PPTX audit pack:\n" + "\n".join(offenders)
